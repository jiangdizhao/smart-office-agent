from __future__ import annotations

from typing import Any, Literal

from app.models import ToolResult
from app.presentation_config import presentation_config
from app.tools.presentation_controller import get_presentation_status

InsightMode = Literal["summary", "explanation"]
Language = Literal["zh", "en"]


def _clean(value: Any, maximum: int = 360) -> str:
    return " ".join(str(value or "").strip().split())[:maximum]


def _shape_content(slide: Any) -> tuple[str, list[str], list[str], list[str], str]:
    title = ""
    title_shape = getattr(slide.shapes, "title", None)
    if title_shape is not None and getattr(title_shape, "has_text_frame", False):
        title = _clean(getattr(title_shape, "text", ""), 220)

    text_blocks: list[str] = []
    table_blocks: list[str] = []
    chart_blocks: list[str] = []
    for shape in slide.shapes:
        if shape is title_shape:
            continue
        if getattr(shape, "has_text_frame", False):
            text = _clean(getattr(shape, "text", ""))
            if text and text not in text_blocks:
                text_blocks.append(text)
        if getattr(shape, "has_table", False):
            rows: list[str] = []
            try:
                for row in shape.table.rows:
                    cells = [_clean(cell.text, 100) for cell in row.cells]
                    row_text = " | ".join(value for value in cells if value)
                    if row_text:
                        rows.append(row_text)
            except Exception:
                rows = []
            if rows:
                table_blocks.append("; ".join(rows[:8]))
        if getattr(shape, "has_chart", False):
            try:
                chart = shape.chart
                chart_title = ""
                if getattr(chart, "has_title", False):
                    chart_title = _clean(chart.chart_title.text_frame.text, 180)
                series_names = [
                    _clean(getattr(series, "name", ""), 80)
                    for series in chart.series
                    if _clean(getattr(series, "name", ""), 80)
                ]
                description = chart_title or " / ".join(series_names)
                if description:
                    chart_blocks.append(description)
            except Exception:
                pass

    notes = ""
    try:
        notes = _clean(slide.notes_slide.notes_text_frame.text, 800)
    except Exception:
        notes = ""
    return title, text_blocks, table_blocks, chart_blocks, notes


def _summary_text(
    *,
    language: Language,
    slide_number: int,
    title: str,
    text_blocks: list[str],
    table_blocks: list[str],
    chart_blocks: list[str],
) -> str:
    topic = title or (f"第 {slide_number} 页" if language == "zh" else f"slide {slide_number}")
    points = [*text_blocks, *table_blocks, *chart_blocks][:3]
    if language == "en":
        if not points:
            return f"Slide {slide_number} is titled “{topic}”. No readable body text, table or chart label was found."
        return f"Slide {slide_number} focuses on “{topic}”. Key points: {'; '.join(points)}."
    if not points:
        return f"第 {slide_number} 页的主题是“{topic}”。当前没有提取到可读的正文、表格或图表标签。"
    return f"第 {slide_number} 页的主题是“{topic}”。关键内容包括：{'；'.join(points)}。"


def _explanation_text(
    *,
    language: Language,
    slide_number: int,
    title: str,
    text_blocks: list[str],
    table_blocks: list[str],
    chart_blocks: list[str],
    notes: str,
) -> str:
    topic = title or (f"第 {slide_number} 页" if language == "zh" else f"slide {slide_number}")
    evidence = [*text_blocks, *table_blocks, *chart_blocks]
    first = evidence[0] if evidence else ""
    second = evidence[1] if len(evidence) > 1 else ""
    supporting = notes or (evidence[-1] if evidence else "")
    if language == "en":
        if not evidence:
            return f"Slide {slide_number}, “{topic}”, appears to rely mainly on visual content. I cannot infer its intended message reliably from text alone."
        logic = f"It starts with {first}"
        if second:
            logic += f", then develops the point through {second}"
        meaning = supporting or first
        return f"Slide {slide_number} is explaining “{topic}”. {logic}. The main takeaway is {meaning}."
    if not evidence:
        return f"第 {slide_number} 页“{topic}”主要依赖视觉内容，仅凭文本还不能可靠判断它想表达的完整含义。"
    logic = f"它先提出“{first}”"
    if second:
        logic += f"，再通过“{second}”继续展开"
    meaning = supporting or first
    return f"第 {slide_number} 页主要在解释“{topic}”。{logic}。这页希望听众记住的核心是“{meaning}”。"


def current_slide_insight(
    *,
    mode: InsightMode,
    language: Language = "zh",
) -> ToolResult:
    tool_name = (
        "presentation_summarize_current_slide"
        if mode == "summary"
        else "presentation_explain_current_slide"
    )
    source_path = presentation_config.presentation_path.resolve()
    if not source_path.is_file():
        return ToolResult(
            tool_name=tool_name,
            ok=False,
            message=f"Configured presentation was not found: {source_path}",
            data={
                "execution_mode": "failed",
                "requested_state": {"current_slide_insight": True, "mode": mode},
                "presentation_path": str(source_path),
            },
        )

    try:
        from pptx import Presentation

        status_result = get_presentation_status()
        status = dict(status_result.data)
        current = status.get("current_slide")
        slide_number = int(current) if isinstance(current, int) and current >= 1 else 1
        deck = Presentation(str(source_path))
        if slide_number > len(deck.slides):
            slide_number = len(deck.slides)
        if slide_number < 1:
            raise RuntimeError("The configured presentation has no slides.")
        slide = deck.slides[slide_number - 1]
        title, text_blocks, table_blocks, chart_blocks, notes = _shape_content(slide)
        spoken = (
            _summary_text(
                language=language,
                slide_number=slide_number,
                title=title,
                text_blocks=text_blocks,
                table_blocks=table_blocks,
                chart_blocks=chart_blocks,
            )
            if mode == "summary"
            else _explanation_text(
                language=language,
                slide_number=slide_number,
                title=title,
                text_blocks=text_blocks,
                table_blocks=table_blocks,
                chart_blocks=chart_blocks,
                notes=notes,
            )
        )
        return ToolResult(
            tool_name=tool_name,
            ok=True,
            message=spoken,
            artifacts=[str(source_path)],
            data={
                "execution_mode": "real",
                "requested_state": {"current_slide_insight": True, "mode": mode},
                "current_slide_insight": True,
                "insight_mode": mode,
                "language": language,
                "slide_number": slide_number,
                "slide_title": title,
                "slide_text": text_blocks,
                "slide_tables": table_blocks,
                "slide_charts": chart_blocks,
                "slide_notes": notes,
                "spoken_insight": spoken,
                "presentation_status": status,
                "presentation_path": str(source_path),
            },
        )
    except Exception as exc:
        return ToolResult(
            tool_name=tool_name,
            ok=False,
            message=f"Current slide could not be read: {exc}",
            data={
                "execution_mode": "failed",
                "requested_state": {"current_slide_insight": True, "mode": mode},
                "error": str(exc),
            },
        )
