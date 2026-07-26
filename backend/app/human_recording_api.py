from __future__ import annotations

import asyncio
import hashlib
import json
import os
import re
from datetime import UTC, datetime
from pathlib import Path
from typing import Any, Literal

import aiofiles
from docx import Document
from docx.shared import Pt
from fastapi import APIRouter, File, Form, HTTPException, UploadFile
from fastapi.responses import FileResponse
from pydantic import BaseModel

from app.openai_services import (
    generate_response_text,
    parse_json_object,
    transcribe_human_conversation,
)
from app.presentation_config import presentation_config

router = APIRouter(tags=["human-conversation-recording"])

Language = Literal["zh", "en"]
_AUDIO_PREFIX = "human_conversation_"
_SUMMARY_PREFIX = "human_conversation_summary_"
_ALLOWED_AUDIO_SUFFIXES = {".webm", ".m4a", ".mp4", ".wav", ".mp3", ".ogg"}
_MAX_UPLOAD_BYTES = 250 * 1024 * 1024


class RecordingSummaryRequest(BaseModel):
    language: Language = "zh"


class RecordingSummaryResponse(BaseModel):
    ok: bool
    route: Literal["human_recording_summary"] = "human_recording_summary"
    spoken_text: str
    recording_available: bool
    artifact_url: str | None = None
    document_path: str | None = None
    transcript_path: str | None = None
    audio_path: str | None = None
    opened_in_word: bool = False
    transcription_model: str | None = None
    summary_model: str | None = None


def _output_directory() -> Path:
    directory = presentation_config.output_directory.resolve()
    directory.mkdir(parents=True, exist_ok=True)
    return directory


def _conversation_token(conversation_id: str) -> str:
    return hashlib.sha256(conversation_id.encode("utf-8", errors="ignore")).hexdigest()[:16]


def _timestamp() -> str:
    return datetime.now(UTC).astimezone().strftime("%Y%m%d_%H%M%S_%f")


def _safe_audio_suffix(filename: str | None, content_type: str | None) -> str:
    suffix = Path(filename or "").suffix.casefold()
    if suffix in _ALLOWED_AUDIO_SUFFIXES:
        return suffix
    mapping = {
        "audio/webm": ".webm",
        "audio/mp4": ".m4a",
        "audio/mpeg": ".mp3",
        "audio/wav": ".wav",
        "audio/x-wav": ".wav",
        "audio/ogg": ".ogg",
    }
    return mapping.get((content_type or "").split(";", 1)[0].casefold(), ".webm")


def _latest_audio(conversation_id: str) -> Path | None:
    token = _conversation_token(conversation_id)
    candidates = sorted(
        (
            path
            for path in _output_directory().glob(f"{_AUDIO_PREFIX}{token}_*")
            if path.suffix.casefold() in _ALLOWED_AUDIO_SUFFIXES
        ),
        key=lambda path: path.stat().st_mtime,
        reverse=True,
    )
    return candidates[0].resolve() if candidates else None


def _speaker_transcript(transcription: dict[str, Any]) -> str:
    segments = transcription.get("segments")
    if isinstance(segments, list) and segments:
        lines: list[str] = []
        for segment in segments:
            if not isinstance(segment, dict):
                continue
            speaker = str(segment.get("speaker") or "Speaker")
            text = str(segment.get("text") or "").strip()
            if not text:
                continue
            start = float(segment.get("start") or 0.0)
            minutes = int(start // 60)
            seconds = int(start % 60)
            lines.append(f"[{minutes:02d}:{seconds:02d}] Speaker {speaker}: {text}")
        if lines:
            return "\n".join(lines)
    return str(transcription.get("text") or "").strip()


def _summary_instructions(language: Language) -> str:
    if language == "en":
        return """
You summarize a recording of a conversation between people in a room. The Smart Office Agent was not a participant.
Use only the supplied transcript. Preserve uncertainty caused by unclear audio or speaker diarization.
Return one JSON object and no surrounding prose with these exact keys:
"title" string,
"overview" string,
"key_points" array of strings,
"decisions" array of strings,
"action_items" array of objects with "owner", "task", and "deadline" strings,
"open_questions" array of strings.
Use clear professional English. Do not invent names, decisions, deadlines, or action items.
""".strip()
    return """
你正在总结一段现场人员之间的对话录音，Smart Office Agent 不是对话参与者。
只能依据提供的转写文本；录音不清或说话人区分不确定时，必须保留这种不确定性。
仅返回一个 JSON 对象，不得添加其他文字。JSON 必须使用以下键：
"title" 字符串，
"overview" 字符串，
"key_points" 字符串数组，
"decisions" 字符串数组，
"action_items" 对象数组，每个对象包含 "owner"、"task"、"deadline" 三个字符串，
"open_questions" 字符串数组。
使用清晰、专业的中文。不得编造姓名、决定、期限或行动事项。
""".strip()


def _list_value(payload: dict[str, Any], key: str) -> list[Any]:
    value = payload.get(key)
    return value if isinstance(value, list) else []


def _text_value(payload: dict[str, Any], key: str, fallback: str) -> str:
    value = payload.get(key)
    return str(value).strip() if value is not None and str(value).strip() else fallback


def _write_summary_docx(
    *,
    document_path: Path,
    summary: dict[str, Any],
    transcript: str,
    audio_path: Path,
    language: Language,
    transcription_model: str,
    summary_model: str,
) -> None:
    document = Document()
    styles = document.styles
    styles["Normal"].font.name = "Microsoft YaHei" if language == "zh" else "Aptos"
    styles["Normal"].font.size = Pt(10.5)

    default_title = "现场对话总结" if language == "zh" else "Human Conversation Summary"
    document.add_heading(_text_value(summary, "title", default_title), level=0)
    generated = datetime.now(UTC).astimezone().isoformat(timespec="seconds")
    metadata = document.add_paragraph()
    metadata.add_run(
        ("生成时间：" if language == "zh" else "Generated: ") + generated
    ).bold = True
    metadata.add_run("\n")
    metadata.add_run(
        ("源录音：" if language == "zh" else "Source recording: ") + audio_path.name
    )
    metadata.add_run("\n")
    metadata.add_run(
        ("转写模型：" if language == "zh" else "Transcription model: ")
        + transcription_model
    )
    metadata.add_run("\n")
    metadata.add_run(
        ("总结模型：" if language == "zh" else "Summary model: ") + summary_model
    )

    document.add_heading("总体概述" if language == "zh" else "Overview", level=1)
    document.add_paragraph(
        _text_value(
            summary,
            "overview",
            "未生成总体概述。" if language == "zh" else "No overview was generated.",
        )
    )

    sections = [
        ("要点" if language == "zh" else "Key points", "key_points"),
        ("已作决定" if language == "zh" else "Decisions", "decisions"),
        ("待解决问题" if language == "zh" else "Open questions", "open_questions"),
    ]
    for heading, key in sections:
        document.add_heading(heading, level=1)
        values = [str(item).strip() for item in _list_value(summary, key) if str(item).strip()]
        if values:
            for value in values:
                document.add_paragraph(value, style="List Bullet")
        else:
            document.add_paragraph("无。" if language == "zh" else "None recorded.")

    document.add_heading("行动事项" if language == "zh" else "Action items", level=1)
    actions = [item for item in _list_value(summary, "action_items") if isinstance(item, dict)]
    if actions:
        table = document.add_table(rows=1, cols=3)
        table.style = "Table Grid"
        headers = (
            ("负责人", "任务", "期限")
            if language == "zh"
            else ("Owner", "Task", "Deadline")
        )
        for cell, header in zip(table.rows[0].cells, headers, strict=True):
            cell.text = header
        for action in actions:
            cells = table.add_row().cells
            cells[0].text = str(action.get("owner") or ("未指定" if language == "zh" else "Unassigned"))
            cells[1].text = str(action.get("task") or "")
            cells[2].text = str(action.get("deadline") or ("未指定" if language == "zh" else "Not specified"))
    else:
        document.add_paragraph("无。" if language == "zh" else "None recorded.")

    document.add_page_break()
    document.add_heading("完整转写" if language == "zh" else "Full transcript", level=1)
    document.add_paragraph(transcript or ("没有可用转写。" if language == "zh" else "No transcript was available."))
    document.save(document_path)


def _open_in_word(path: Path) -> bool:
    if os.name != "nt":
        return False
    try:
        os.startfile(str(path))  # type: ignore[attr-defined]
        return True
    except OSError:
        return False


def _artifact_url(path: Path) -> str:
    return f"/api/human-recordings/artifacts/{path.name}"


@router.post("/api/human-recordings/{conversation_id}")
async def upload_human_recording(
    conversation_id: str,
    file: UploadFile = File(...),
    language: Language = Form("zh"),
) -> dict[str, Any]:
    suffix = _safe_audio_suffix(file.filename, file.content_type)
    token = _conversation_token(conversation_id)
    path = (_output_directory() / f"{_AUDIO_PREFIX}{token}_{_timestamp()}{suffix}").resolve()
    total = 0
    try:
        async with aiofiles.open(path, "wb") as output:
            while chunk := await file.read(1024 * 1024):
                total += len(chunk)
                if total > _MAX_UPLOAD_BYTES:
                    raise HTTPException(status_code=413, detail="Recording is larger than 250 MB.")
                await output.write(chunk)
    except Exception:
        path.unlink(missing_ok=True)
        raise
    finally:
        await file.close()

    if total == 0:
        path.unlink(missing_ok=True)
        raise HTTPException(status_code=400, detail="The uploaded recording is empty.")

    metadata = {
        "conversation_id": conversation_id,
        "language": language,
        "audio_path": str(path),
        "audio_filename": path.name,
        "content_type": file.content_type,
        "size_bytes": total,
        "uploaded_at": datetime.now(UTC).isoformat(),
    }
    metadata_path = path.with_suffix(path.suffix + ".json")
    metadata_path.write_text(json.dumps(metadata, ensure_ascii=False, indent=2), encoding="utf-8")
    return {
        "ok": True,
        "recording_available": True,
        "audio_path": str(path),
        "audio_filename": path.name,
        "artifact_url": _artifact_url(path),
        "size_bytes": total,
    }


@router.post(
    "/api/human-recordings/{conversation_id}/summary",
    response_model=RecordingSummaryResponse,
)
async def summarize_human_recording(
    conversation_id: str,
    req: RecordingSummaryRequest,
) -> RecordingSummaryResponse:
    audio_path = _latest_audio(conversation_id)
    if audio_path is None:
        spoken = (
            "无法总结，因为当前没有已停止并保存的现场对话录音。请先点击开始录音，结束后再点击停止录音。"
            if req.language == "zh"
            else "I cannot summarize because there is no saved human-conversation recording. Start recording and stop it before requesting a summary."
        )
        return RecordingSummaryResponse(
            ok=False,
            spoken_text=spoken,
            recording_available=False,
        )

    try:
        transcription = await transcribe_human_conversation(
            audio_path,
            language=req.language,
        )
        transcript = _speaker_transcript(transcription)
        if not transcript.strip():
            raise RuntimeError("The recording transcription was empty.")

        max_chars = max(10_000, int(os.getenv("SMART_OFFICE_SUMMARY_TRANSCRIPT_MAX_CHARS", "120000")))
        model_input = transcript[:max_chars]
        if len(transcript) > max_chars:
            model_input += "\n[Transcript truncated by configured character limit.]"

        summary_text, summary_model = await generate_response_text(
            input_text=model_input,
            instructions=_summary_instructions(req.language),
            max_output_tokens=2600,
        )
        try:
            summary_payload = parse_json_object(summary_text)
        except Exception:
            summary_payload = {
                "title": "现场对话总结" if req.language == "zh" else "Human Conversation Summary",
                "overview": summary_text,
                "key_points": [],
                "decisions": [],
                "action_items": [],
                "open_questions": [],
            }

        token = _conversation_token(conversation_id)
        stamp = _timestamp()
        output_directory = _output_directory()
        transcript_path = output_directory / f"{_SUMMARY_PREFIX}{token}_{stamp}_transcript.txt"
        transcript_json_path = output_directory / f"{_SUMMARY_PREFIX}{token}_{stamp}_transcript.json"
        document_path = output_directory / f"{_SUMMARY_PREFIX}{token}_{stamp}.docx"
        transcript_path.write_text(transcript, encoding="utf-8")
        transcript_json_path.write_text(
            json.dumps(transcription, ensure_ascii=False, indent=2, default=str),
            encoding="utf-8",
        )
        await asyncio.to_thread(
            _write_summary_docx,
            document_path=document_path,
            summary=summary_payload,
            transcript=transcript,
            audio_path=audio_path,
            language=req.language,
            transcription_model=str(transcription.get("model") or "unknown"),
            summary_model=summary_model,
        )
        opened = await asyncio.to_thread(_open_in_word, document_path)
    except Exception as exc:
        raise HTTPException(status_code=502, detail=f"Human conversation summary failed: {exc}") from exc

    spoken = (
        "现场人员对话已经完成转写和总结，DOCX 总结稿已经创建并在 Word 中打开。"
        if req.language == "zh" and opened
        else "现场人员对话已经完成转写和总结，DOCX 总结稿已经创建。请打开生成的文件。"
        if req.language == "zh"
        else "The human conversation was transcribed and summarized. The DOCX summary was created and opened in Word."
        if opened
        else "The human conversation was transcribed and summarized. The DOCX summary was created; please open the generated file."
    )
    return RecordingSummaryResponse(
        ok=True,
        spoken_text=spoken,
        recording_available=True,
        artifact_url=_artifact_url(document_path),
        document_path=str(document_path),
        transcript_path=str(transcript_path),
        audio_path=str(audio_path),
        opened_in_word=opened,
        transcription_model=str(transcription.get("model") or "unknown"),
        summary_model=summary_model,
    )


@router.get("/api/human-recordings/artifacts/{filename}")
def human_recording_artifact(filename: str):
    safe_name = Path(filename).name
    if safe_name != filename:
        raise HTTPException(status_code=400, detail="Invalid artifact filename.")
    if not safe_name.startswith((_AUDIO_PREFIX, _SUMMARY_PREFIX)):
        raise HTTPException(status_code=404, detail="Artifact not found.")
    if not re.fullmatch(r"[A-Za-z0-9_.-]+", safe_name):
        raise HTTPException(status_code=404, detail="Artifact not found.")

    output_directory = _output_directory()
    path = (output_directory / safe_name).resolve()
    if output_directory not in path.parents or not path.is_file():
        raise HTTPException(status_code=404, detail="Artifact not found.")

    suffix = path.suffix.casefold()
    media_types = {
        ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ".json": "application/json",
        ".txt": "text/plain",
        ".webm": "audio/webm",
        ".m4a": "audio/mp4",
        ".mp4": "audio/mp4",
        ".wav": "audio/wav",
        ".mp3": "audio/mpeg",
        ".ogg": "audio/ogg",
    }
    media_type = media_types.get(suffix)
    if media_type is None:
        raise HTTPException(status_code=404, detail="Artifact not found.")
    return FileResponse(path, media_type=media_type, filename=path.name)
